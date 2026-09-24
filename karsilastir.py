import re
from collections import Counter
from datetime import datetime

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


MAC_DESENI = re.compile(r"^([0-9A-F]{2}:){5}[0-9A-F]{2}$")
GECERSIZ_MACLER = {"00:00:00:00:00:00", "FF:FF:FF:FF:FF:FF"}

IZLENEN_ALANLAR = ["IP Adresi", "MAC Adresi", "Hostname", "Uretici", "Acik Portlar"]


def rapor_oku(dosya_yolu):
    """Excel raporunu okuyup {ip: cihaz_sozlugu} yapisina cevirir."""
    kitap = load_workbook(dosya_yolu, read_only=True)
    sayfa = kitap.active

    satirlar = list(sayfa.iter_rows(values_only=True))
    if not satirlar:
        kitap.close()
        return {}

    basliklar = [str(b).strip() if b else "" for b in satirlar[0]]
    cihazlar = {}

    for satir in satirlar[1:]:
        if not satir or not satir[0]:
            continue
        kayit = dict(zip(basliklar, satir))
        ip = str(kayit.get("IP Adresi", "")).strip()
        if ip:
            cihazlar[ip] = kayit

    kitap.close()
    return cihazlar


def mac_normallestir(mac):
    """'08-00-27-aa-bb-cc' -> '08:00:27:AA:BB:CC'. Gecersizse None doner."""
    if mac is None:
        return None
    metin = str(mac).strip().upper().replace("-", ":")
    if MAC_DESENI.match(metin) and metin not in GECERSIZ_MACLER:
        return metin
    return None


def _ip(kayit):
    return str(kayit.get("IP Adresi", "")).strip()


def _mac(kayit):
    return mac_normallestir(kayit.get("MAC Adresi"))


def _ip_sirasi(ip):
    try:
        return tuple(int(p) for p in str(ip).split("."))
    except ValueError:
        return (999, str(ip))


def anahtarla(cihazlar):
    """{ip: kayit} yapisini {anahtar: kayit} yapisina cevirir.

    Gecerli ve raporda tek olan MAC varsa anahtar MAC'tir.
    MAC bilinmiyorsa ya da ayni MAC birden fazla IP'de goruluyorsa IP kullanilir.
    """
    sayac = Counter(_mac(k) for k in cihazlar.values())
    sonuc = {}
    for ip, kayit in cihazlar.items():
        mac = _mac(kayit)
        if mac and sayac[mac] == 1:
            sonuc[f"MAC:{mac}"] = kayit
        else:
            sonuc[f"IP:{ip}"] = kayit
    return sonuc


def _alan_farklari(eski, yeni):
    farklar = []
    for alan in IZLENEN_ALANLAR:
        eski_deger = str(eski.get(alan) if eski.get(alan) is not None else "-").strip()
        yeni_deger = str(yeni.get(alan) if yeni.get(alan) is not None else "-").strip()

        if alan == "MAC Adresi":
            ayni = (mac_normallestir(eski_deger) or eski_deger) == (mac_normallestir(yeni_deger) or yeni_deger)
        else:
            ayni = eski_deger == yeni_deger

        if not ayni:
            farklar.append(f"{alan}: '{eski_deger}' -> '{yeni_deger}'")
    return farklar


def karsilastir(eski_dosya, yeni_dosya):
    """Iki raporu kiyaslar, farklari dondurur.

    Eslestirme once MAC adresine gore yapilir; boylece DHCP ile IP'si degisen
    cihaz 'kaybolan + yeni' yerine 'degisen' olarak raporlanir.
    """
    eski = anahtarla(rapor_oku(eski_dosya))
    yeni = anahtarla(rapor_oku(yeni_dosya))

    ortak = eski.keys() & yeni.keys()
    eslesmeler = [(eski[a], yeni[a]) for a in ortak]
    eski_kalan = {a: eski[a] for a in eski.keys() - ortak}
    yeni_kalan = {a: yeni[a] for a in yeni.keys() - ortak}

    # 2. tur: MAC'i taraflardan birinde bilinmeyen cihazlari IP ile eslestir.
    # Iki tarafta da MAC biliniyor ve farkliysa bunlar farkli cihazdir, eslestirilmez.
    yeni_ip_dizini = {_ip(k): a for a, k in yeni_kalan.items()}
    for a_eski, k_eski in list(eski_kalan.items()):
        a_yeni = yeni_ip_dizini.get(_ip(k_eski))
        if a_yeni is None or a_yeni not in yeni_kalan:
            continue
        k_yeni = yeni_kalan[a_yeni]
        if _mac(k_eski) and _mac(k_yeni):
            continue
        eslesmeler.append((k_eski, k_yeni))
        del eski_kalan[a_eski]
        del yeni_kalan[a_yeni]

    yeni_cihazlar = sorted(yeni_kalan.values(), key=lambda k: _ip_sirasi(_ip(k)))
    kaybolanlar = sorted(eski_kalan.values(), key=lambda k: _ip_sirasi(_ip(k)))

    degisenler = []
    for k_eski, k_yeni in sorted(eslesmeler, key=lambda e: _ip_sirasi(_ip(e[1]))):
        farklar = _alan_farklari(k_eski, k_yeni)
        if farklar:
            degisenler.append({"ip": _ip(k_yeni), "farklar": farklar})

    return {
        "yeni": yeni_cihazlar,
        "kaybolan": kaybolanlar,
        "degisen": degisenler,
    }


def fark_raporu_yaz(sonuc, dosya_adi=None):
    """Karsilastirma sonucunu Excel dosyasina yazar."""
    if dosya_adi is None:
        damga = datetime.now().strftime("%Y%m%d_%H%M")
        dosya_adi = f"fark_{damga}.xlsx"

    kitap = Workbook()
    kitap.remove(kitap.active)

    basliklar = ["IP Adresi", "Hostname", "MAC Adresi", "Uretici", "Acik Portlar"]

    def sayfa_olustur(ad, kayitlar, renk):
        sayfa = kitap.create_sheet(ad)
        for sutun, baslik in enumerate(basliklar, start=1):
            hucre = sayfa.cell(row=1, column=sutun, value=baslik)
            hucre.font = Font(bold=True, color="FFFFFF")
            hucre.fill = PatternFill("solid", fgColor=renk)
            hucre.alignment = Alignment(horizontal="center")

        for satir, kayit in enumerate(kayitlar, start=2):
            for sutun, baslik in enumerate(basliklar, start=1):
                sayfa.cell(row=satir, column=sutun, value=kayit.get(baslik, "-"))

        for sutun in range(1, len(basliklar) + 1):
            sayfa.column_dimensions[get_column_letter(sutun)].width = 22
        sayfa.freeze_panes = "A2"

    sayfa_olustur("Yeni Cihazlar", sonuc["yeni"], "2E7D32")
    sayfa_olustur("Kaybolan Cihazlar", sonuc["kaybolan"], "C62828")

    sayfa = kitap.create_sheet("Degisenler")
    for sutun, baslik in enumerate(["IP Adresi", "Degisiklik"], start=1):
        hucre = sayfa.cell(row=1, column=sutun, value=baslik)
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = PatternFill("solid", fgColor="EF6C00")
        hucre.alignment = Alignment(horizontal="center")

    satir = 2
    for kayit in sonuc["degisen"]:
        for fark in kayit["farklar"]:
            sayfa.cell(row=satir, column=1, value=kayit["ip"])
            sayfa.cell(row=satir, column=2, value=fark)
            satir += 1

    sayfa.column_dimensions["A"].width = 18
    sayfa.column_dimensions["B"].width = 70
    sayfa.freeze_panes = "A2"

    kitap.save(dosya_adi)
    return dosya_adi


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 3:
        print("Kullanim: python karsilastir.py <eski_rapor.xlsx> <yeni_rapor.xlsx>")
        sys.exit(1)

    sonuc = karsilastir(sys.argv[1], sys.argv[2])

    print(f"\n[+] Yeni cihaz    : {len(sonuc['yeni'])}")
    for c in sonuc["yeni"]:
        print(f"      {c.get('IP Adresi')}  {c.get('MAC Adresi', '-')}  {c.get('Uretici', '-')}")

    print(f"\n[-] Kaybolan cihaz: {len(sonuc['kaybolan'])}")
    for c in sonuc["kaybolan"]:
        print(f"      {c.get('IP Adresi')}  {c.get('MAC Adresi', '-')}  {c.get('Uretici', '-')}")

    print(f"\n[!] Degisen cihaz : {len(sonuc['degisen'])}")
    for c in sonuc["degisen"]:
        print(f"      {c['ip']}")
        for f in c["farklar"]:
            print(f"        {f}")

    dosya = fark_raporu_yaz(sonuc)
    print(f"\n[+] Fark raporu: {dosya}")