from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime

from cihaz_tipi import tip_tahmin


# (Excel basligi, cihaz sozlugundeki anahtar)
# karsilastir.py sutunlari basliga gore okur; sira degisse de eski raporlarla kiyas bozulmaz.
SUTUNLAR = [
    ("IP Adresi", "ip"),
    ("Tip", "tip"),
    ("Hostname", "hostname"),
    ("MAC Adresi", "mac"),
    ("Uretici", "uretici"),
    ("Acik Portlar", "portlar"),
    ("SNMP Adi", "snmp"),
    ("Sarf Malzeme", "toner"),
    ("Tip Gerekcesi", "tip_gerekce"),
]

BASLIKLAR = [baslik for baslik, _ in SUTUNLAR]


def _tip_ekle(cihaz):
    """Cihazda tip yoksa tahmin edip ekler (orijinal sozlugu degistirmez)."""
    if cihaz.get("tip"):
        return cihaz
    tahmin = tip_tahmin(cihaz)
    return {**cihaz,
            "tip": tahmin["tip"],
            "tip_gerekce": f"{tahmin['guven']} guven: {tahmin['gerekce']}"}


def excel_yaz(cihazlar, dosya_adi=None):
    """Cihaz listesini bicimlendirilmis bir Excel dosyasina yazar."""
    if dosya_adi is None:
        damga = datetime.now().strftime("%Y%m%d_%H%M")
        dosya_adi = f"envanter_{damga}.xlsx"

    kitap = Workbook()
    sayfa = kitap.active
    sayfa.title = "Ag Envanteri"

    baslik_yazi = Font(bold=True, color="FFFFFF")
    baslik_dolgu = PatternFill("solid", fgColor="2F5597")

    for sutun, baslik in enumerate(BASLIKLAR, start=1):
        hucre = sayfa.cell(row=1, column=sutun, value=baslik)
        hucre.font = baslik_yazi
        hucre.fill = baslik_dolgu
        hucre.alignment = Alignment(horizontal="center")

    for satir, cihaz in enumerate(cihazlar, start=2):
        cihaz = _tip_ekle(cihaz)
        for sutun, (_, anahtar) in enumerate(SUTUNLAR, start=1):
            sayfa.cell(row=satir, column=sutun, value=cihaz.get(anahtar, "-"))

    for sutun in range(1, len(BASLIKLAR) + 1):
        en_uzun = max(
            (len(str(sayfa.cell(row=s, column=sutun).value or ""))
             for s in range(1, sayfa.max_row + 1)),
            default=10
        )
        sayfa.column_dimensions[get_column_letter(sutun)].width = min(en_uzun + 4, 50)

    sayfa.freeze_panes = "A2"
    kitap.save(dosya_adi)
    return dosya_adi


if __name__ == "__main__":
    ornek = [
        {"ip": "192.168.1.1", "hostname": "modem.local", "mac": "aa:bb:cc:dd:ee:ff",
         "uretici": "Zyxel", "portlar": "80/HTTP, 443/HTTPS"},
        {"ip": "192.168.1.15", "hostname": "-", "mac": "11:22:33:44:55:66",
         "uretici": "Hikvision", "portlar": "80/HTTP"},
    ]
    yol = excel_yaz(ornek)
    print(f"Rapor olusturuldu: {yol}")